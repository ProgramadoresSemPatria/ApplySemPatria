"""Unit tests for ResumeChameleon keyword extraction and headline merge."""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pytest

SCRIPTS = Path(__file__).resolve().parent.parent / "scripts"
sys.path.insert(0, str(SCRIPTS))

from resume_keywords import (  # noqa: E402
    extract_role_keywords,
    merge_headline_skill_list,
    parse_pipe_skills,
    split_headline_paragraph,
)
from resume_chameleon import (  # noqa: E402
    generate_for_job,
    pick_best_master,
    score_master,
    sync_master_keywords,
)
from resume_pdf import (  # noqa: E402
    DEFAULT_PDF_FONT_PATH,
    compact_sparse_trailing_pages,
    replace_pdf_skill_lines,
)


def test_merge_headline_keywords_requirements_first_no_duplicates():
    existing = ["Python", "RAG", "LangChain", "AWS", "Terraform", "CI/CD", "Tests", "Rest API", "Cloud", "LangGraph"]
    required = ["langfuse", "python", "langsmith", "mcp"]
    merged = merge_headline_skill_list(required, existing)
    assert merged[0].lower() == "langfuse"
    assert "Python" in merged
    assert len(merged) == 13


def test_extract_role_keywords_from_job_text():
    text = "Senior AI Engineer — LangGraph, LangChain, Python, AWS, RAG required."
    keywords = extract_role_keywords(text, lexicon=["langgraph", "langchain", "python", "aws", "rag", "kotlin"])
    assert "langgraph" in [k.lower() for k in keywords]
    assert "kotlin" not in [k.lower() for k in keywords]


def test_split_headline_paragraph():
    text = "Senior AI Engineer | Senior Software Engineer\nPython | RAG | LangChain"
    title, skills = split_headline_paragraph(text)
    assert "Senior AI Engineer" in title
    assert skills.startswith("Python")


def test_pick_best_master_by_keyword_overlap(tmp_path: Path):
    ai_path = tmp_path / "ai.docx"
    android_path = tmp_path / "android.docx"
    ai_path.write_bytes(b"x")
    android_path.write_bytes(b"x")
    cfg = {
        "headline_separator": " | ",
        "headline_paragraph_index": 1,
        "masters": [
            {"id": "ai", "label": "AI", "path": str(ai_path), "keywords": ["python", "langchain", "rag"], "default": False},
            {"id": "android", "label": "Android", "path": str(android_path), "keywords": ["kotlin", "android"], "default": True},
        ],
    }
    role_keywords = ["langchain", "python", "aws"]
    assert score_master(cfg["masters"][0], role_keywords, cfg) >= score_master(cfg["masters"][1], role_keywords, cfg)
    best = pick_best_master(cfg, role_keywords)
    assert best is not None
    assert best["id"] == "ai"


@pytest.fixture
def sample_master_docx(tmp_path: Path) -> Path:
    from docx import Document

    path = tmp_path / "master.docx"
    doc = Document()
    doc.add_paragraph("JANE DOE")
    doc.add_paragraph(
        "Senior AI Engineer\nPython | RAG | LangChain | AWS | Terraform | CI/CD | Tests | Rest API"
    )
    doc.add_paragraph("Summary")
    doc.save(str(path))
    return path


def test_generate_for_job_docx_edits_headline(sample_master_docx: Path, tmp_path: Path, monkeypatch):
    cfg = {
        "enabled": True,
        "output_format": "docx",
        "download_dir": str(tmp_path / "downloads"),
        "cache_dir": str(tmp_path / "cache"),
        "headline_separator": " | ",
        "headline_max_chars": 320,
        "headline_paragraph_index": 1,
        "tech_lexicon": ["langfuse", "langsmith", "mcp"],
        "masters": [
            {
                "id": "ai",
                "label": "AI",
                "path": str(sample_master_docx),
                "format": "docx",
                "default": True,
                "keywords": [],
            }
        ],
    }

    monkeypatch.setattr("resume_chameleon.load_chameleon_config", lambda track_id=None: cfg)
    monkeypatch.setattr("resume_chameleon.chameleon_is_configured", lambda c=None: True)

    job = {
        "source": "linkedin_posts",
        "role": "AI Engineer",
        "company": "Acme AI",
        "description_snippet": "Need Langfuse, LangSmith, MCP, and Python experience.",
        "url": "https://www.linkedin.com/posts/test-chameleon",
    }

    result = generate_for_job(job, track_id="ai-engineer", download=True)
    assert result["headline_edited"] is True
    assert "Langfuse" in result["headline_after"] or "langfuse" in result["headline_after"].lower()
    out = Path(result["output_path"])
    assert out.is_file()
    meta_path = Path(result["cache_path"]).with_suffix(".json")
    # meta lives beside cache docx with .json extension via artifact_paths
    cache_json = list((tmp_path / "cache" / "ai-engineer").glob("*.json"))[0]
    meta = json.loads(cache_json.read_text())
    assert meta["master_id"] == "ai"


def test_sync_master_keywords_reads_docx(sample_master_docx: Path, tmp_path: Path, monkeypatch):
    cfg = {
        "headline_separator": " | ",
        "headline_paragraph_index": 1,
        "masters": [{"id": "ai", "label": "AI", "path": str(sample_master_docx), "format": "docx", "keywords": []}],
    }
    monkeypatch.setattr("resume_chameleon.load_chameleon_config", lambda track_id=None: cfg)
    monkeypatch.setattr("resume_chameleon.save_chameleon_config", lambda track_id, data: None)
    updated = sync_master_keywords("ai-engineer", save=False)
    keywords = updated["masters"][0]["keywords"]
    assert "Python" in keywords
    assert "LangChain" in keywords


@pytest.fixture
def sample_master_pdf(tmp_path: Path) -> Path:
    pymupdf = pytest.importorskip("pymupdf")
    fitz = pymupdf
    font_path = Path(DEFAULT_PDF_FONT_PATH)
    if not font_path.is_file():
        pytest.skip("Tahoma font unavailable")

    path = tmp_path / "master.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 50), "JANE DOE", fontsize=16)
    page.insert_text(
        (72, 78),
        "Senior AI Engineer | Senior Software Engineer",
        fontname="Tahoma",
        fontfile=str(font_path),
        fontsize=11.5,
        color=(0.43, 0.47, 0.47),
    )
    page.insert_text(
        (72, 94),
        "Python | RAG | LangChain | AWS | Terraform | CI/CD",
        fontname="Tahoma",
        fontfile=str(font_path),
        fontsize=11.5,
        color=(0.43, 0.47, 0.47),
    )
    page.insert_text((72, 120), "Experience and summary content", fontsize=10)
    doc.new_page()
    doc[1].insert_text((72, 72), "More experience", fontsize=10)
    doc.save(str(path))
    doc.close()
    return path


def test_replace_pdf_skill_lines_keeps_two_page_master(sample_master_pdf: Path, tmp_path: Path):
    dest = tmp_path / "tailored.pdf"
    cfg = {"pdf_font_path": DEFAULT_PDF_FONT_PATH}
    layout = replace_pdf_skill_lines(
        sample_master_pdf,
        dest,
        ["Langfuse | Python | RAG", "LangChain | AWS | Terraform"],
        cfg=cfg,
    )
    pymupdf = pytest.importorskip("pymupdf")
    assert pymupdf.open(str(dest)).page_count == 2
    assert layout["master_pages"] == 2
    assert layout["output_pages"] == 2


def test_compact_sparse_trailing_pages_merges_short_tail(tmp_path: Path):
    pymupdf = pytest.importorskip("pymupdf")
    fitz = pymupdf
    path = tmp_path / "sparse-tail.pdf"
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Page 1", fontsize=12)
    doc.new_page().insert_text((72, 72), "Page 2 content", fontsize=10)
    for i in range(40):
        doc[1].insert_text((72, 120 + i * 16), f"Experience line {i}", fontsize=10)
    doc.new_page().insert_text((72, 72), "Education", fontsize=12)
    doc[2].insert_text((72, 100), "University 2016-2020", fontsize=10)
    doc.save(str(path))
    doc.close()

    doc = fitz.open(str(path))
    merged = compact_sparse_trailing_pages(
        doc,
        cfg={
            "compact_trailing_pages": True,
            "pdf_bottom_margin": 20.0,
        },
    )
    assert merged == 1
    assert doc.page_count == 2
    assert "Education" in doc[1].get_text()
    doc.close()
