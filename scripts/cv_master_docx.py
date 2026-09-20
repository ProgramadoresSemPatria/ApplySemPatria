#!/usr/bin/env python3
"""Parse and build CV master DOCX files using the MASTER_CV template layout."""

from __future__ import annotations

import re
import shutil
from pathlib import Path
from typing import Any

from cv_master_schema import (
    CvAchievement,
    CvContact,
    CvEducation,
    CvExperience,
    CvProfile,
    CvProject,
)

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_TEMPLATE = ROOT / "templates" / "cv-master" / "MASTER_CV.docx"

HEADLINE_PARAGRAPH_INDEX = 1
SUMMARY_BODY_FALLBACK_INDEX = 6

SECTION_SUMMARY = "Summary"
SECTION_ACHIEVEMENTS = "Key Achievements"
SECTION_EXPERIENCE = "Experience"
SECTION_PROJECTS = "Github projects"
SECTION_EDUCATION = "Education"


def _para_style(paragraph: Any) -> str:
    return paragraph.style.name if paragraph.style else ""


def _is_heading(paragraph: Any, level: int) -> bool:
    style = _para_style(paragraph)
    return style == f"Heading {level}"


def _section_indices(paragraphs: list[Any]) -> dict[str, int]:
    out: dict[str, int] = {}
    for idx, para in enumerate(paragraphs):
        if _is_heading(para, 1):
            title = (para.text or "").strip()
            if title and title not in out:
                out[title] = idx
    return out


def _section_body_paragraphs(paragraphs: list[Any], start: int, end: int) -> list[Any]:
    return paragraphs[start + 1 : end]


def _parse_contact_line(text: str) -> CvContact:
    raw = (text or "").strip()
    contact = CvContact()
    if not raw:
        return contact
    email_match = re.search(r"[\w.+-]+@[\w.-]+\.\w+", raw)
    if email_match:
        contact.email = email_match.group(0)
    phone_match = re.search(r"\+?\d[\d\s().-]{7,}\d", raw)
    if phone_match:
        contact.phone = phone_match.group(0).strip()
    linkedin_match = re.search(r"(https?://[^\s]+linkedin[^\s]+|linkedin\.com/[^\s]+)", raw, re.I)
    if linkedin_match:
        contact.linkedin_url = linkedin_match.group(0)
    portfolio_match = re.search(r"(https?://[^\s]+)", raw)
    if portfolio_match and "linkedin" not in portfolio_match.group(0).lower():
        contact.portfolio_url = portfolio_match.group(0)
    if "Brazil" in raw or "GMT" in raw:
        loc_match = re.search(r"([A-Za-zÀ-ú .,-]+(?:GMT[+-]\d+)?)", raw.split("LinkedIn")[-1])
        if loc_match:
            contact.location = loc_match.group(1).strip()
    return contact


def _parse_experience_block(lines: list[Any]) -> list[CvExperience]:
    jobs: list[CvExperience] = []
    i = 0
    while i < len(lines):
        para = lines[i]
        text = (para.text or "").strip()
        if not text:
            i += 1
            continue
        if _is_heading(para, 2):
            i += 1
            continue
        if "\t" in text and not text.startswith("Stack:"):
            company, location = text.split("\t", 1)
            company = company.strip()
            location = location.strip()
            role = ""
            date_range = ""
            context = ""
            bullets: list[str] = []
            stack = ""
            i += 1
            if i < len(lines) and _is_heading(lines[i], 2):
                role_line = (lines[i].text or "").strip()
                if "\t" in role_line:
                    role, date_range = role_line.split("\t", 1)
                else:
                    role = role_line
                i += 1
            if i < len(lines) and _is_heading(lines[i], 3):
                context = (lines[i].text or "").strip()
                i += 1
            while i < len(lines):
                line_text = (lines[i].text or "").strip()
                if not line_text:
                    i += 1
                    continue
                if _is_heading(lines[i], 2):
                    break
                if "\t" in line_text and not line_text.startswith("Stack:") and not role:
                    break
                if line_text.startswith("Stack:"):
                    stack = line_text.removeprefix("Stack:").strip()
                    i += 1
                    continue
                if _is_heading(lines[i], 3):
                    break
                bullets.append(line_text)
                i += 1
            jobs.append(
                CvExperience(
                    company=company,
                    location=location,
                    role=role.strip(),
                    date_range=date_range.strip(),
                    context=context,
                    bullets=bullets,
                    stack=stack,
                )
            )
            continue
        i += 1
    return jobs


def _parse_achievements(lines: list[Any]) -> list[CvAchievement]:
    items: list[CvAchievement] = []
    i = 0
    while i < len(lines):
        if _is_heading(lines[i], 2):
            title = (lines[i].text or "").strip()
            body_parts: list[str] = []
            i += 1
            while i < len(lines) and not _is_heading(lines[i], 2):
                chunk = (lines[i].text or "").strip()
                if chunk:
                    body_parts.append(chunk)
                i += 1
            items.append(CvAchievement(title=title, body=" ".join(body_parts)))
            continue
        i += 1
    return items


def _parse_projects(lines: list[Any]) -> list[CvProject]:
    projects: list[CvProject] = []
    i = 0
    while i < len(lines):
        text = (lines[i].text or "").strip()
        if not text or _is_heading(lines[i], 1):
            i += 1
            continue
        title = text
        desc = ""
        i += 1
        if i < len(lines):
            nxt = (lines[i].text or "").strip()
            if nxt and not _is_heading(lines[i], 1):
                desc = nxt
                i += 1
        projects.append(CvProject(title=title, description=desc))
    return projects


def _parse_education(lines: list[Any]) -> list[CvEducation]:
    rows: list[CvEducation] = []
    i = 0
    while i < len(lines):
        text = (lines[i].text or "").strip()
        if not text:
            i += 1
            continue
        institution = text
        degree = ""
        date_range = ""
        i += 1
        if i < len(lines):
            line = (lines[i].text or "").strip()
            if line and "\t" in line:
                degree, date_range = line.split("\t", 1)
                i += 1
        rows.append(
            CvEducation(
                institution=institution.strip(),
                degree=degree.strip(),
                date_range=date_range.strip(),
            )
        )
    return rows


def parse_master_docx(path: Path | str) -> CvProfile:
    from docx import Document  # noqa: WPS433

    doc_path = Path(path)
    doc = Document(str(doc_path))
    paragraphs = doc.paragraphs
    sections = _section_indices(paragraphs)

    full_name = (paragraphs[0].text or "").strip() if paragraphs else ""
    headline_line = (paragraphs[HEADLINE_PARAGRAPH_INDEX].text or "").strip() if len(paragraphs) > 1 else ""
    contact = _parse_contact_line(paragraphs[2].text if len(paragraphs) > 2 else "")

    summary = ""
    if SECTION_SUMMARY in sections:
        start = sections[SECTION_SUMMARY]
        end = min(
            [idx for title, idx in sections.items() if idx > start],
            default=len(paragraphs),
        )
        body = _section_body_paragraphs(paragraphs, start, end)
        summary_parts = [(p.text or "").strip() for p in body if (p.text or "").strip()]
        summary = " ".join(summary_parts)
    elif len(paragraphs) > SUMMARY_BODY_FALLBACK_INDEX:
        summary = (paragraphs[SUMMARY_BODY_FALLBACK_INDEX].text or "").strip()

    achievements: list[CvAchievement] = []
    if SECTION_ACHIEVEMENTS in sections:
        start = sections[SECTION_ACHIEVEMENTS]
        end = sections.get(SECTION_EXPERIENCE, len(paragraphs))
        achievements = _parse_achievements(_section_body_paragraphs(paragraphs, start, end))

    experience: list[CvExperience] = []
    if SECTION_EXPERIENCE in sections:
        start = sections[SECTION_EXPERIENCE]
        lang_indices = [idx for title, idx in sections.items() if title == "Languages" and idx > start]
        end = min(lang_indices) if lang_indices else len(paragraphs)
        experience = _parse_experience_block(_section_body_paragraphs(paragraphs, start, end))

    programming_languages = ""
    spoken_languages = ""
    lang_starts = sorted(idx for title, idx in sections.items() if title == "Languages")
    if lang_starts:
        first = lang_starts[0]
        second_end = lang_starts[1] if len(lang_starts) > 1 else sections.get(SECTION_PROJECTS, len(paragraphs))
        prog_body = _section_body_paragraphs(paragraphs, first, second_end)
        for para in prog_body:
            text = (para.text or "").strip()
            if text.startswith("•") or "years" in text.lower():
                programming_languages = text
                break
    if len(lang_starts) > 1:
        start = lang_starts[-1]
        spoken_body = _section_body_paragraphs(paragraphs, start, len(paragraphs))
        spoken_parts = [(p.text or "").strip() for p in spoken_body if (p.text or "").strip()]
        spoken_languages = " ".join(spoken_parts)

    projects: list[CvProject] = []
    if SECTION_PROJECTS in sections:
        start = sections[SECTION_PROJECTS]
        end = sections.get(SECTION_EDUCATION, len(paragraphs))
        projects = _parse_projects(_section_body_paragraphs(paragraphs, start, end))

    education: list[CvEducation] = []
    if SECTION_EDUCATION in sections:
        start = sections[SECTION_EDUCATION]
        end = lang_starts[-1] if lang_starts else len(paragraphs)
        education = _parse_education(_section_body_paragraphs(paragraphs, start, end))

    return CvProfile(
        full_name=full_name,
        headline_line=headline_line,
        contact=contact,
        summary=summary,
        achievements=achievements,
        experience=experience,
        programming_languages=programming_languages,
        projects=projects,
        education=education,
        spoken_languages=spoken_languages,
    )


def _format_contact_line(contact: CvContact) -> str:
    parts: list[str] = []
    if contact.phone.strip():
        parts.append(contact.phone.strip())
    if contact.email.strip():
        parts.append(contact.email.strip())
    if contact.linkedin_url.strip():
        parts.append("LinkedIn" if contact.linkedin_url.startswith("http") else contact.linkedin_url)
    if contact.portfolio_url.strip():
        parts.append("Portfolio")
    if contact.location.strip():
        parts.append(contact.location.strip())
    return "  ".join(parts)


def _delete_paragraph(paragraph: Any) -> None:
    element = paragraph._element  # noqa: SLF001
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def _insert_paragraph_after(paragraph: Any, text: str = "", style: str | None = None) -> Any:
    from docx.oxml import OxmlElement  # noqa: WPS433
    from docx.text.paragraph import Paragraph  # noqa: WPS433

    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)  # noqa: SLF001
    new_para = Paragraph(new_p, paragraph._parent)
    if style:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def _rebuild_section(doc: Any, heading_title: str, builder) -> None:
    """Replace body content under a Heading 1 section."""
    paragraphs = doc.paragraphs
    sections = _section_indices(paragraphs)
    if heading_title not in sections:
        return
    start = sections[heading_title]
    heading_para = paragraphs[start]
    next_starts = [idx for title, idx in sections.items() if idx > start]
    end = min(next_starts) if next_starts else len(paragraphs)
    to_delete = paragraphs[start + 1 : end]
    for para in reversed(to_delete):
        _delete_paragraph(para)
    anchor = heading_para
    for chunk in builder():
        if isinstance(chunk, tuple):
            text, style = chunk
            anchor = _insert_paragraph_after(anchor, text, style)
        else:
            anchor = _insert_paragraph_after(anchor, chunk, "normal")


def build_master_docx(
    profile: CvProfile | dict[str, Any],
    dest: Path | str,
    *,
    template: Path | str | None = None,
) -> Path:
    from docx import Document  # noqa: WPS433

    prof = profile if isinstance(profile, CvProfile) else CvProfile.from_dict(profile)
    template_path = Path(template or DEFAULT_TEMPLATE)
    if not template_path.is_file():
        raise FileNotFoundError(f"CV template not found: {template_path}")

    out = Path(dest)
    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, out)
    doc = Document(str(out))

    if doc.paragraphs:
        doc.paragraphs[0].text = prof.full_name.strip()
    if len(doc.paragraphs) > HEADLINE_PARAGRAPH_INDEX:
        doc.paragraphs[HEADLINE_PARAGRAPH_INDEX].text = prof.headline_line.strip()
    if len(doc.paragraphs) > 2:
        doc.paragraphs[2].text = _format_contact_line(prof.contact)

    def summary_builder():
        if prof.summary.strip():
            yield prof.summary.strip()

    _rebuild_section(doc, SECTION_SUMMARY, summary_builder)

    def achievements_builder():
        for item in prof.achievements:
            yield (item.title, "Heading 2")
            if item.body.strip():
                yield item.body.strip()

    _rebuild_section(doc, SECTION_ACHIEVEMENTS, achievements_builder)

    def experience_builder():
        for job in prof.experience:
            company_line = f"   {job.company.strip()}\t{job.location.strip()}".rstrip()
            yield company_line
            role_line = f"{job.role.strip()}\t{job.date_range.strip()}".strip()
            yield (role_line, "Heading 2")
            if job.context.strip():
                yield (job.context.strip(), "Heading 3")
            for bullet in job.bullets:
                if bullet.strip():
                    yield bullet.strip()
            if job.stack.strip():
                yield f"Stack: {job.stack.strip()}"

    _rebuild_section(doc, SECTION_EXPERIENCE, experience_builder)

    def projects_builder():
        for project in prof.projects:
            yield project.title.strip()
            if project.description.strip():
                yield project.description.strip()

    _rebuild_section(doc, SECTION_PROJECTS, projects_builder)

    def education_builder():
        for row in prof.education:
            yield row.institution.strip()
            if row.degree.strip() or row.date_range.strip():
                yield f"{row.degree.strip()}\t{row.date_range.strip()}".strip("\t")

    _rebuild_section(doc, SECTION_EDUCATION, education_builder)

    doc.save(str(out))
    return out


def default_template_path() -> Path:
    return DEFAULT_TEMPLATE
